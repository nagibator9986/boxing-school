"""Возврат клиентов, этап 1: разбор Excel, сверка с ботом, проверка исхода в CRM.

Номера и имена здесь вымышленные: репозиторий публичный, а настоящий список — это
переписки живых родителей.
"""

from __future__ import annotations

import io
import sqlite3
import subprocess
from contextlib import closing
import sys
from datetime import UTC, datetime, timedelta
from pathlib import Path

import pytest
from openpyxl import Workbook

from app.admin.admin_store import AdminStore
from app.campaign.contacts import (
    BotFacts,
    CampaignOutcome,
    CampaignSegment,
    CampaignStatus,
    apply_outcome,
    initial_decision,
    reconcile,
    segment_for,
)
from crm.campaign_excel import parse_workbook
from crm.config import CrmConfig
from tests.test_crm import client, sandbox  # noqa: F401 - фикстуры CRM

ROOT = Path(__file__).resolve().parent.parent
NOW = datetime(2026, 10, 22, 6, 0, tzinfo=UTC)

HEADER = [
    "Номер WhatsApp", "Имя из WhatsApp", "Имя родителя из текста", "Имя ребёнка / участника",
    "Этап и статус", "Сценарий продолжения", "Краткий итог переписки", "Почему включён / исключён",
    "Препятствие со слов клиента", "Следующий шаг", "Нужна проверка", "Что проверить",
    "Последнее сообщение (UTC+5)", "Без общения, дней", "Пробное / курс", "Оплата по переписке",
    "Связаться не раньше (UTC+5)", "Фрагменты переписки", "Статус первого анализа",
    "Был в первом отборе", "Метод повторного анализа", "Ключ контакта",
]


def _row(phone: str, stage: str, *, check: str = "Да", parent: str = "", first: str = "",
         summary: str = "Спрашивал про бокс для сына", last: datetime | None = None,
         later: datetime | None = None, fragments: str = "Родитель: здравствуйте…") -> list[object]:
    values: dict[str, object] = {
        "Номер WhatsApp": phone, "Имя из WhatsApp": "Профиль", "Имя родителя из текста": parent,
        "Этап и статус": stage, "Краткий итог переписки": summary, "Нужна проверка": check,
        "Что проверить": "Был ли на пробном", "Последнее сообщение (UTC+5)": last or datetime(2026, 10, 2, 15, 30),
        "Связаться не раньше (UTC+5)": later, "Фрагменты переписки": fragments,
        "Статус первого анализа": first, "Ключ контакта": f"key-{phone[-4:]}",
    }
    return [values.get(name) for name in HEADER]


def _xlsx(rows: list[list[object]], *, sheet: str = "Потенциальные") -> bytes:
    """Файл того же вида, что даёт разбор: сводка сверху, заголовок в восьмой строке."""
    book = Workbook()
    page = book.active
    page.title = sheet
    page.append(["Потенциальные клиенты — повторный анализ"])
    for _ in range(6):
        page.append([])
    page.append(HEADER)
    for row in rows:
        page.append(row)
    book.create_sheet("Все контакты").append(["другой лист"])
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


# --------------------------------------------------------------------------- #
# Правила
# --------------------------------------------------------------------------- #
@pytest.mark.parametrize(
    ("stage", "first", "segment"),
    [
        ("Интерес затих до записи", "", CampaignSegment.FADED),
        ("Интерес затих до записи", "Компания не вернулась к запросу", CampaignSegment.NO_REPLY),
        ("Записан — исход пробного неизвестен", "", CampaignSegment.BOOKED_UNKNOWN),
        ("Записан - исход пробного неизвестен", "", CampaignSegment.BOOKED_UNKNOWN),
        ("После пробного / бесплатного курса", "", CampaignSegment.AFTER_TRIAL),
        ("Запросил счёт — оплата неизвестна", "", CampaignSegment.INVOICE),
        ("Запросил счет – оплата неизвестна", "", CampaignSegment.INVOICE),
        ("Просил вернуться позже", "", CampaignSegment.LATER),
        ("Бывший посетитель — возможен возврат", "", CampaignSegment.FORMER),
        ("Текущий диалог", "", CampaignSegment.LIVE),
        ("Другие занятия — проверить предложение", "", CampaignSegment.OTHER_PROGRAM),
        ("Действующий клиент / сообщил об оплате", "", CampaignSegment.CURRENT_CLIENT),
        ("Что-то новое", "", CampaignSegment.UNKNOWN),
    ],
)
def test_stage_maps_to_segment(stage: str, first: str, segment: CampaignSegment) -> None:
    assert segment_for(stage, first) is segment


def test_only_a_clean_faded_contact_is_ready_at_once() -> None:
    assert initial_decision(CampaignSegment.FADED, check_required=False).status is CampaignStatus.READY
    assert initial_decision(CampaignSegment.NO_REPLY, check_required=False).status is CampaignStatus.READY
    assert initial_decision(CampaignSegment.FADED, check_required=True).status is CampaignStatus.NEEDS_CHECK
    # Был на пробном или просил счёт — неизвестно главное: занимается ли уже.
    for segment in (CampaignSegment.AFTER_TRIAL, CampaignSegment.BOOKED_UNKNOWN, CampaignSegment.INVOICE,
                    CampaignSegment.LATER, CampaignSegment.FORMER, CampaignSegment.UNKNOWN):
        assert initial_decision(segment, check_required=False).status is CampaignStatus.NEEDS_CHECK, segment
    for segment in (CampaignSegment.LIVE, CampaignSegment.OTHER_PROGRAM, CampaignSegment.CURRENT_CLIENT,
                    CampaignSegment.REFUSED, CampaignSegment.NOT_CLIENT):
        decision = initial_decision(segment, check_required=False)
        assert decision.status is CampaignStatus.EXCLUDED and decision.reason, segment


CHECKED = NOW - timedelta(days=4)


def _reconcile(facts: BotFacts, status: CampaignStatus = CampaignStatus.READY, *, checked_at=None):
    return reconcile(status, CampaignSegment.FADED, facts, now=NOW, checked_at=checked_at)


def test_hard_bot_signals_act_even_after_a_check() -> None:
    assert _reconcile(BotFacts(ignored=True), checked_at=CHECKED).status is CampaignStatus.EXCLUDED
    assert _reconcile(BotFacts(stop_word=True), checked_at=CHECKED).status is CampaignStatus.OPTED_OUT
    assert _reconcile(BotFacts(lead_status="converted"), checked_at=CHECKED).status is CampaignStatus.EXCLUDED


def test_a_trial_ahead_is_a_wait_not_a_permanent_exclusion() -> None:
    """Запись впереди — не «не пишем» навсегда: после пробного нужен вопрос «пришёл ли»."""
    ahead = _reconcile(BotFacts(lead_status="trial_booked", trial_at=NOW + timedelta(days=2)), checked_at=CHECKED)
    assert (ahead.status, ahead.segment) == (CampaignStatus.NEEDS_CHECK, CampaignSegment.BOOKED_UNKNOWN)
    assert "после него" in ahead.reason


def test_soft_bot_signals_send_the_contact_back_to_a_check() -> None:
    past = _reconcile(BotFacts(lead_status="trial_booked", trial_at=NOW - timedelta(days=5)))
    assert (past.status, past.segment) == (CampaignStatus.NEEDS_CHECK, CampaignSegment.BOOKED_UNKNOWN)
    assert _reconcile(BotFacts(human_pause=True)).status is CampaignStatus.NEEDS_CHECK
    fresh = _reconcile(BotFacts(last_activity_at=NOW - timedelta(days=3)))
    assert fresh.status is CampaignStatus.NEEDS_CHECK and "свежий" in fresh.reason
    assert _reconcile(BotFacts(last_activity_at=NOW - timedelta(days=30))) is None


def test_a_check_stands_until_the_client_writes_again() -> None:
    before = NOW - timedelta(days=5)
    # Всё это человек видел, когда отмечал: решение остаётся.
    assert _reconcile(BotFacts(last_activity_at=before, last_inbound_at=before), checked_at=CHECKED) is None
    assert _reconcile(
        BotFacts(lead_status="trial_booked", trial_at=before, last_activity_at=before, last_inbound_at=before),
        checked_at=CHECKED,
    ) is None
    # Напоминание бота после отметки — не новость.
    assert _reconcile(BotFacts(last_activity_at=NOW - timedelta(days=1), last_inbound_at=before),
                      checked_at=CHECKED) is None
    # Клиент сам написал боту после отметки — проверить заново, даже когда разговор уже стих.
    wrote = NOW - timedelta(days=2)
    fresh = _reconcile(BotFacts(last_activity_at=wrote, last_inbound_at=wrote), checked_at=CHECKED)
    assert fresh.status is CampaignStatus.NEEDS_CHECK
    quiet = reconcile(CampaignStatus.READY, CampaignSegment.FADED,
                      BotFacts(last_activity_at=wrote, last_inbound_at=wrote),
                      now=NOW + timedelta(days=30), checked_at=CHECKED)
    assert quiet.status is CampaignStatus.NEEDS_CHECK and "после проверки" in quiet.reason


def test_a_contact_already_written_to_is_left_alone() -> None:
    for status in (CampaignStatus.SENT, CampaignStatus.REPLIED, CampaignStatus.OPTED_OUT, CampaignStatus.EXCLUDED):
        assert _reconcile(BotFacts(stop_word=True, ignored=True), status) is None, status


def test_outcomes() -> None:
    assert apply_outcome(CampaignOutcome.BOUGHT, CampaignSegment.AFTER_TRIAL).status is CampaignStatus.EXCLUDED
    came = apply_outcome(CampaignOutcome.CAME, CampaignSegment.BOOKED_UNKNOWN)
    assert (came.status, came.segment) == (CampaignStatus.READY, CampaignSegment.AFTER_TRIAL)
    no_show = apply_outcome(CampaignOutcome.NO_SHOW, CampaignSegment.BOOKED_UNKNOWN)
    assert (no_show.status, no_show.segment) == (CampaignStatus.READY, CampaignSegment.NO_SHOW)
    assert apply_outcome(CampaignOutcome.WRITE, CampaignSegment.FADED).segment is CampaignSegment.FADED
    assert apply_outcome(CampaignOutcome.REFUSED, CampaignSegment.FADED).status is CampaignStatus.EXCLUDED
    assert apply_outcome(CampaignOutcome.NOT_CLIENT, CampaignSegment.FADED).status is CampaignStatus.EXCLUDED
    with pytest.raises(ValueError):
        apply_outcome(CampaignOutcome.LATER, CampaignSegment.LATER)
    later = apply_outcome(CampaignOutcome.LATER, CampaignSegment.LATER, not_before=NOW)
    assert later.status is CampaignStatus.READY


# --------------------------------------------------------------------------- #
# Excel
# --------------------------------------------------------------------------- #
def test_the_workbook_is_read_by_column_names() -> None:
    parsed = parse_workbook(io.BytesIO(_xlsx([
        _row("+7 (701) 555-01-01", "Интерес затих до записи", check="Нет", parent="Петрова Анна"),
        _row("87015550102", "После пробного / бесплатного курса"),
        _row("7015550103", "Интерес затих до записи", first="Компания не вернулась к запросу"),
        _row("77015550101", "Текущий диалог"),          # повтор номера — пропускается
        _row("12345", "Интерес затих до записи"),        # не номер
        _row("77015550104", "Просил вернуться позже", later=datetime(2026, 11, 13)),
        _row("77015550105", "Интерес затих до записи", check=""),
    ])))

    assert parsed.sheet == "Потенциальные"
    assert [row.chat_id for row in parsed.rows] == [
        "77015550101", "77015550102", "77015550103", "77015550104", "77015550105"
    ]
    assert parsed.duplicates == 1
    assert len(parsed.errors) == 1 and "Строка 13" in parsed.errors[0], parsed.errors
    first = parsed.rows[0]
    assert (first.segment, first.check_required, first.parent_name) == (CampaignSegment.FADED, False, "Петрова Анна")
    assert parsed.rows[1].check_required is True
    # Пустая пометка — сомнение: проверяет человек.
    assert parsed.rows[4].check_required is True
    assert parsed.rows[2].segment is CampaignSegment.NO_REPLY
    # Время в файле — по Костанаю (UTC+5).
    assert first.last_contact_at == datetime(2026, 10, 2, 10, 30, tzinfo=UTC)
    assert parsed.rows[3].not_before == datetime(2026, 11, 12, 19, 0, tzinfo=UTC)
    # Цитаты из переписок не читаются вовсе.
    assert not hasattr(first, "fragments")


def test_a_wrong_file_is_explained_not_crashed() -> None:
    book = Workbook()
    book.active.append(["Имя", "Телефон"])
    buffer = io.BytesIO()
    book.save(buffer)
    assert "нет листа" in parse_workbook(io.BytesIO(buffer.getvalue())).errors[0]
    assert "не читается" in parse_workbook(io.BytesIO(b"not an excel file")).errors[0]


# --------------------------------------------------------------------------- #
# CRM
# --------------------------------------------------------------------------- #
def _upload(web, rows: list[list[object]]):
    return web.post(
        "/campaign/import",
        data={"file": (io.BytesIO(_xlsx(rows)), "разбор.xlsx")},
        content_type="multipart/form-data",
        follow_redirects=True,
    )


def _status(cfg: CrmConfig, chat_id: str) -> tuple[str, str, str | None]:
    with closing(sqlite3.connect(cfg.bot_db)) as conn, conn:
        return conn.execute(
            "SELECT status, segment, outcome FROM campaign_contact WHERE chat_id = ?", (chat_id,)
        ).fetchone()


def _contact_id(cfg: CrmConfig, chat_id: str) -> str:
    with closing(sqlite3.connect(cfg.bot_db)) as conn, conn:
        return conn.execute("SELECT id FROM campaign_contact WHERE chat_id = ?", (chat_id,)).fetchone()[0]


def _bot_conversation(cfg: CrmConfig, conv_id: str, chat_id: str, *, last_inbound: datetime,
                      blocked: bool = False, lead: str | None = None, trial: datetime | None = None,
                      pause: str | None = None) -> None:
    stamp = last_inbound.strftime("%Y-%m-%d %H:%M:%S")
    with closing(sqlite3.connect(cfg.bot_db)) as conn, conn:
        conn.execute(
            "INSERT INTO conversation (id, conv_key, channel_id, chat_type, chat_id, lang, lang_locked, state,"
            " msg_in_count, msg_out_count, bot_miss_count, followup_stage, followup_blocked, created_at,"
            " updated_at, last_inbound_at) VALUES (?, ?, 'wa', 'whatsapp', ?, 'ru', 0, 'active', 1, 1, 0, 0, ?,"
            " ?, ?, ?)",
            (conv_id, f"{conv_id}:whatsapp:{chat_id}", chat_id, 1 if blocked else 0, stamp, stamp, stamp),
        )
        if lead:
            conn.execute(
                "INSERT INTO lead (id, conversation_id, created_at, updated_at, channel, child_name, child_age,"
                " child_gender, phone_source, status, escalation, messages_count, trial_slot)"
                " VALUES (?, ?, ?, ?, 'whatsapp', 'Тимур', 8, 'male', 'channel', ?, 0, 2, ?)",
                (f"l-{conv_id}", conv_id, stamp, stamp, lead, trial.strftime("%Y-%m-%d %H:%M:%S") if trial else None),
            )
        if pause:
            conn.execute(
                "INSERT INTO escalation_state (conversation_id, paused, pause_reason, escalation_count,"
                " resume_policy) VALUES (?, 1, ?, 0, 'manual_only')",
                (conv_id, pause),
            )


def test_upload_sorts_the_list_into_tabs(client, sandbox: CrmConfig) -> None:  # noqa: F811
    page = _upload(client, [
        _row("77015550101", "Интерес затих до записи", check="Нет", parent="Петрова Анна"),
        _row("77015550102", "После пробного / бесплатного курса"),
        _row("77015550103", "Текущий диалог"),
    ]).get_data(as_text=True)

    assert "новых 3" in page, page
    assert _status(sandbox, "77015550101")[0] == "ready"
    assert _status(sandbox, "77015550102")[0] == "needs_check"
    assert _status(sandbox, "77015550103")[0] == "excluded"
    check_tab = client.get("/campaign/?tab=check").get_data(as_text=True)
    assert "+7 701 555 01 02" in check_tab and "Петрова Анна" not in check_tab
    assert "Петрова Анна" in client.get("/campaign/?tab=ready").get_data(as_text=True)
    assert "Возврат клиентов" in client.get("/").get_data(as_text=True)


def test_the_bot_knows_what_the_file_does_not(client, sandbox: CrmConfig) -> None:  # noqa: F811
    now = datetime.now(tz=UTC)
    _bot_conversation(sandbox, "cv1", "77015550201", last_inbound=now - timedelta(days=20),
                      lead="trial_booked", trial=now + timedelta(days=1))
    _bot_conversation(sandbox, "cv2", "77015550202", last_inbound=now - timedelta(days=20), blocked=True)
    _bot_conversation(sandbox, "cv3", "77015550203", last_inbound=now - timedelta(days=2))
    _bot_conversation(sandbox, "cv4", "77015550204", last_inbound=now - timedelta(days=30),
                      lead="trial_booked", trial=now - timedelta(days=20))
    _bot_conversation(sandbox, "cv7", "77015550207", last_inbound=now - timedelta(days=20), pause="operator_reply")
    _bot_conversation(sandbox, "cv8", "77015550208", last_inbound=now - timedelta(days=20), pause="escalation")
    AdminStore(sandbox.admin_db).set("ignored_numbers", "+7 701 555 02 05")

    _upload(client, [_row(f"7701555020{n}", "Интерес затих до записи", check="Нет") for n in range(1, 9)])

    assert _status(sandbox, "77015550201")[:2] == ("needs_check", "booked_unknown")  # пробное впереди
    assert _status(sandbox, "77015550202")[0] == "opted_out"      # стоп-слово боту
    assert _status(sandbox, "77015550203")[0] == "needs_check"    # свежий диалог
    assert _status(sandbox, "77015550204")[:2] == ("needs_check", "booked_unknown")
    assert _status(sandbox, "77015550205")[0] == "excluded"       # сотрудник
    assert _status(sandbox, "77015550206")[0] == "ready"          # бот его не знает
    assert _status(sandbox, "77015550207")[0] == "needs_check"    # диалог ведёт администратор
    assert _status(sandbox, "77015550208")[0] == "ready"          # пауза самого бота — не человек
    with closing(sqlite3.connect(sandbox.bot_db)) as conn, conn:
        linked = conn.execute(
            "SELECT conversation_id FROM campaign_contact WHERE chat_id = '77015550203'"
        ).fetchone()[0]
    assert linked == "cv3"


def test_a_check_mark_moves_the_contact_and_survives_a_new_upload(client, sandbox: CrmConfig) -> None:  # noqa: F811
    rows = [_row("77015550301", "Записан — исход пробного неизвестен"),
            _row("77015550302", "После пробного / бесплатного курса")]
    _upload(client, rows)

    page = client.post(
        f"/campaign/{_contact_id(sandbox, '77015550301')}/outcome",
        data={"outcome": "no_show", "back": "list", "tab": "check"},
        follow_redirects=True,
    )
    assert "Отмечено: Не пришёл" in page.get_data(as_text=True)
    assert page.request.path == "/campaign/"
    assert _status(sandbox, "77015550301") == ("ready", "no_show", "no_show")
    client.post(f"/campaign/{_contact_id(sandbox, '77015550302')}/outcome", data={"outcome": "bought"})
    assert _status(sandbox, "77015550302") == ("excluded", "after_trial", "bought")

    rows[1] = _row("77015550302", "После пробного / бесплатного курса", summary="Новый итог")
    report = _upload(client, rows).get_data(as_text=True)

    assert "без смены статуса (проверены вами, просили не писать или им уже писали) 2" in report, report
    assert _status(sandbox, "77015550302") == ("excluded", "after_trial", "bought")
    with closing(sqlite3.connect(sandbox.bot_db)) as conn, conn:
        assert conn.execute(
            "SELECT summary FROM campaign_contact WHERE chat_id = '77015550302'"
        ).fetchone()[0] == "Новый итог"


def test_a_new_file_recounts_what_only_the_file_decided(client, sandbox: CrmConfig) -> None:  # noqa: F811
    """«Текущий диалог» прошлого разбора в новом стал «интерес затих» — контакт снова к работе."""
    _upload(client, [_row("77015551501", "Текущий диалог")])
    assert _status(sandbox, "77015551501")[0] == "excluded"

    _upload(client, [_row("77015551501", "Интерес затих до записи", check="Нет")])

    assert _status(sandbox, "77015551501")[:2] == ("ready", "faded")


def test_later_needs_a_date_and_keeps_it(client, sandbox: CrmConfig) -> None:  # noqa: F811
    _upload(client, [_row("77015550401", "Просил вернуться позже", later=datetime(2026, 11, 13))])
    contact = _contact_id(sandbox, "77015550401")

    page = client.post(f"/campaign/{contact}/outcome", data={"outcome": "later"}, follow_redirects=True)
    assert "укажите дату" in page.get_data(as_text=True)
    assert _status(sandbox, "77015550401")[0] == "needs_check"

    client.post(f"/campaign/{contact}/outcome", data={"outcome": "later", "not_before": "2026-12-01"})
    with closing(sqlite3.connect(sandbox.bot_db)) as conn, conn:
        status, not_before = conn.execute(
            "SELECT status, not_before FROM campaign_contact WHERE id = ?", (contact,)
        ).fetchone()
    # Начало 1 декабря по Костанаю.
    assert (status, not_before[:16]) == ("ready", "2026-11-30 19:00")

    # «Писать можно» дату, о которой клиент просил сам, не стирает.
    client.post(f"/campaign/{contact}/outcome", data={"outcome": "write"})
    with closing(sqlite3.connect(sandbox.bot_db)) as conn, conn:
        assert conn.execute("SELECT not_before FROM campaign_contact WHERE id = ?", (contact,)).fetchone()[0]


def test_a_check_mark_does_not_override_a_new_stop_word(client, sandbox: CrmConfig) -> None:  # noqa: F811
    """Клиент написал боту «стоп» уже после загрузки списка — «писать можно» его не отменяет."""
    _bot_conversation(sandbox, "cv9", "77015550901", last_inbound=datetime.now(tz=UTC) - timedelta(days=2))
    _upload(client, [_row("77015550901", "Интерес затих до записи", check="Нет")])
    assert _status(sandbox, "77015550901")[0] == "needs_check"
    with closing(sqlite3.connect(sandbox.bot_db)) as conn, conn:
        conn.execute("UPDATE conversation SET followup_blocked = 1 WHERE id = 'cv9'")

    client.post(f"/campaign/{_contact_id(sandbox, '77015550901')}/outcome", data={"outcome": "write"})

    assert _status(sandbox, "77015550901")[0] == "opted_out"


def test_every_dialog_of_a_number_counts(client, sandbox: CrmConfig) -> None:  # noqa: F811
    """Стоп-слово и запись из старого диалога номера действуют, хотя есть диалог новее."""
    now = datetime.now(tz=UTC)
    _bot_conversation(sandbox, "old1", "77015551001", last_inbound=now - timedelta(days=60), blocked=True)
    _bot_conversation(sandbox, "new1", "77015551001", last_inbound=now - timedelta(days=20))
    _bot_conversation(sandbox, "old2", "77015551002", last_inbound=now - timedelta(days=40),
                      lead="trial_booked", trial=now + timedelta(days=3))
    _bot_conversation(sandbox, "new2", "77015551002", last_inbound=now - timedelta(days=20))

    _upload(client, [_row("77015551001", "Интерес затих до записи", check="Нет"),
                     _row("77015551002", "Интерес затих до записи", check="Нет")])

    assert _status(sandbox, "77015551001")[0] == "opted_out"
    assert _status(sandbox, "77015551002")[:2] == ("needs_check", "booked_unknown")
    with closing(sqlite3.connect(sandbox.bot_db)) as conn, conn:
        assert conn.execute(
            "SELECT conversation_id FROM campaign_contact WHERE chat_id = '77015551001'"
        ).fetchone()[0] == "new1"


def test_a_client_who_writes_after_the_check_is_checked_again(client, sandbox: CrmConfig) -> None:  # noqa: F811
    now = datetime.now(tz=UTC)
    _bot_conversation(sandbox, "cv11", "77015551101", last_inbound=now - timedelta(days=30))
    _upload(client, [_row("77015551101", "Интерес затих до записи")])
    client.post(f"/campaign/{_contact_id(sandbox, '77015551101')}/outcome", data={"outcome": "write"})
    assert _status(sandbox, "77015551101")[0] == "ready"

    later = (now + timedelta(minutes=5)).strftime("%Y-%m-%d %H:%M:%S")
    with closing(sqlite3.connect(sandbox.bot_db)) as conn, conn:
        conn.execute("UPDATE conversation SET last_inbound_at = ? WHERE id = 'cv11'", (later,))
    client.post("/campaign/reconcile")

    assert _status(sandbox, "77015551101")[0] == "needs_check"


def test_a_mark_overridden_by_the_bot_is_explained(client, sandbox: CrmConfig) -> None:  # noqa: F811
    now = datetime.now(tz=UTC)
    _bot_conversation(sandbox, "cv12", "77015551201", last_inbound=now - timedelta(days=30),
                      lead="trial_booked", trial=now + timedelta(days=2))
    _upload(client, [_row("77015551201", "Интерес затих до записи")])

    page = client.post(
        f"/campaign/{_contact_id(sandbox, '77015551201')}/outcome", data={"outcome": "write"}, follow_redirects=True
    ).get_data(as_text=True)

    assert "Но по переписке бота контакт теперь «Проверить»" in page
    assert "после него проверить" in page
    assert _status(sandbox, "77015551201")[0] == "needs_check"


def test_a_busy_database_is_a_message_not_an_error_page(client, sandbox: CrmConfig, monkeypatch) -> None:  # noqa: F811
    from crm.campaign_store import CampaignStore

    _upload(client, [_row("77015551301", "Интерес затих до записи")])
    contact = _contact_id(sandbox, "77015551301")
    original = CampaignStore._connect

    def locked(self, *, write: bool = False):
        if write:
            raise sqlite3.OperationalError("database is locked")
        return original(self, write=write)

    monkeypatch.setattr(CampaignStore, "_connect", locked)
    for response in (
        client.post(f"/campaign/{contact}/outcome", data={"outcome": "write"}, follow_redirects=True),
        client.post(f"/campaign/{contact}/restore", follow_redirects=True),
        client.post("/campaign/reconcile", follow_redirects=True),
        _upload(client, [_row("77015551302", "Интерес затих до записи")]),
    ):
        assert response.status_code == 200
        assert "База бота сейчас занята" in response.get_data(as_text=True)
    monkeypatch.undo()
    assert _status(sandbox, "77015551301")[0] == "needs_check"


def test_rows_that_failed_are_named_in_the_report(client) -> None:  # noqa: F811
    page = _upload(client, [_row("77015551401", "Интерес затих до записи"), _row("123", "Интерес затих до записи")])

    assert "Не разобрано строк: 1. Строка 10" in page.get_data(as_text=True)


def test_restore_returns_a_mistake_but_not_a_stop_word(client, sandbox: CrmConfig) -> None:  # noqa: F811
    _bot_conversation(sandbox, "cv5", "77015550502", last_inbound=datetime.now(tz=UTC) - timedelta(days=40),
                      blocked=True)
    _upload(client, [_row("77015550501", "Интерес затих до записи", check="Нет"),
                     _row("77015550502", "Интерес затих до записи", check="Нет")])
    first, second = _contact_id(sandbox, "77015550501"), _contact_id(sandbox, "77015550502")
    client.post(f"/campaign/{first}/outcome", data={"outcome": "refused"})
    assert _status(sandbox, "77015550501")[0] == "excluded"

    client.post(f"/campaign/{first}/restore")
    assert _status(sandbox, "77015550501")[:3] == ("needs_check", "faded", None)

    assert _status(sandbox, "77015550502")[0] == "opted_out"
    page = client.post(f"/campaign/{second}/restore", follow_redirects=True).get_data(as_text=True)
    assert "нельзя" in page and _status(sandbox, "77015550502")[0] == "opted_out"
    # И отметкой «писать можно» стоп-слово не обойти.
    client.post(f"/campaign/{second}/outcome", data={"outcome": "write"})
    assert _status(sandbox, "77015550502")[0] == "opted_out"


def test_the_bots_own_tables_are_only_read(client, sandbox: CrmConfig) -> None:  # noqa: F811
    def snapshot() -> dict[str, list[tuple]]:
        with closing(sqlite3.connect(sandbox.bot_db)) as conn, conn:
            return {
                table: conn.execute(f"SELECT * FROM {table} ORDER BY 1").fetchall()
                for table in ("conversation", "lead", "message", "outbox_message", "escalation_state", "followup_task")
            }

    _bot_conversation(sandbox, "cv6", "77015550601", last_inbound=datetime.now(tz=UTC) - timedelta(days=2))
    before = snapshot()
    _upload(client, [_row("77015550601", "Интерес затих до записи", check="Нет")])
    client.post(f"/campaign/{_contact_id(sandbox, '77015550601')}/outcome", data={"outcome": "write"})
    client.post("/campaign/reconcile")

    assert snapshot() == before


def test_without_the_table_the_section_waits_for_a_restart(client, sandbox: CrmConfig) -> None:  # noqa: F811
    with closing(sqlite3.connect(sandbox.bot_db)) as conn, conn:
        conn.execute("DROP TABLE campaign_contact")

    assert "ещё не создана" in client.get("/campaign/").get_data(as_text=True)
    page = _upload(client, [_row("77015550701", "Интерес затих до записи")]).get_data(as_text=True)
    assert "после перезапуска бота" in page


def test_the_reply_pipeline_does_not_load_the_campaign_code() -> None:
    """Рассылку ведёт свой воркер; конвейер ответов клиенту от её кода не зависит.

    Ответ на сообщение рассылки — обычный разговор с ботом: сломанный код рассылки не
    должен ронять ответы. Код CRM в процесс бота не попадает вовсе.
    """
    probe = (
        "import sys, app.core.pipeline; "
        "print(sorted(name for name in sys.modules if name.startswith(('app.campaign', 'crm.'))))"
    )
    result = subprocess.run([sys.executable, "-c", probe], cwd=ROOT, capture_output=True, text=True, check=True)
    # Последняя строка: выше неё может быть журнал запуска.
    assert result.stdout.strip().splitlines()[-1] == "[]", result.stdout + result.stderr
    probe = "import sys, app.asgi; print(sorted(name for name in sys.modules if name.startswith('crm.campaign')))"
    result = subprocess.run([sys.executable, "-c", probe], cwd=ROOT, capture_output=True, text=True, check=True)
    assert result.stdout.strip().splitlines()[-1] == "[]", result.stdout + result.stderr


# --------------------------------------------------------------------------- #
# Этапы 2–3 в CRM: черновики, одобрение, итог, пульт рассылки
# --------------------------------------------------------------------------- #
def _set(cfg: CrmConfig, chat_id: str, **values: object) -> None:
    names = ", ".join(f"{name} = ?" for name in values)
    with closing(sqlite3.connect(cfg.bot_db)) as conn, conn:
        conn.execute(f"UPDATE campaign_contact SET {names} WHERE chat_id = ?", [*values.values(), chat_id])


def _get(cfg: CrmConfig, chat_id: str, *names: str) -> tuple:
    with closing(sqlite3.connect(cfg.bot_db)) as conn, conn:
        return conn.execute(f"SELECT {', '.join(names)} FROM campaign_contact WHERE chat_id = ?", (chat_id,)).fetchone()


def _draft(client, cfg: CrmConfig, chat_id: str, text: str = "Здравствуйте! Подобрать зал рядом с домом?") -> str:  # noqa: F811
    _upload(client, [_row(chat_id, "Интерес затих до записи", check="Нет")])
    _set(cfg, chat_id, status="draft", draft_text=text, draft_source="model")
    return _contact_id(cfg, chat_id)


def test_a_draft_is_edited_and_approved(client, sandbox: CrmConfig) -> None:  # noqa: F811
    contact = _draft(client, sandbox, "77015550801")
    page = client.get("/campaign/?tab=drafts").get_data(as_text=True)
    assert "Подобрать зал рядом с домом?" in page and "написан по итогу переписки" in page

    page = client.post(f"/campaign/{contact}/message", data={"action": "approve", "text": "Здравствуйте! Свой текст?"},
                       follow_redirects=True).get_data(as_text=True)

    assert "Одобрено" in page
    assert _get(sandbox, "77015550801", "status", "approved_text", "draft_source") == (
        "approved", "Здравствуйте! Свой текст?", "edited",
    )
    assert "Свой текст?" in client.get("/campaign/?tab=approved").get_data(as_text=True)
    client.post(f"/campaign/{contact}/message", data={"action": "unapprove"})
    assert _get(sandbox, "77015550801", "status", "approved_text") == ("draft", None)


def test_approve_all_takes_drafts_as_they_are(client, sandbox: CrmConfig) -> None:  # noqa: F811
    first = _draft(client, sandbox, "77015550811")
    second = _draft(client, sandbox, "77015550812", "Второй текст?")
    _set(sandbox, "77015550812", status="sent")  # уже ушёл — одобрять нечего

    page = client.post("/campaign/approve", data={"ids": [first, second]}, follow_redirects=True).get_data(as_text=True)

    assert "Одобрено: 1" in page and "Не одобрено 1" in page
    assert _get(sandbox, "77015550811", "status", "approved_text")[0] == "approved"
    assert _get(sandbox, "77015550812", "status")[0] == "sent"


def test_approval_does_not_override_a_stop_word(client, sandbox: CrmConfig) -> None:  # noqa: F811
    contact = _draft(client, sandbox, "77015550821")
    _bot_conversation(sandbox, "cv-821", "77015550821", last_inbound=datetime.now(tz=UTC) - timedelta(days=30),
                      blocked=True)

    page = client.post(f"/campaign/{contact}/message", data={"action": "approve", "text": "Текст?"},
                       follow_redirects=True).get_data(as_text=True)

    assert _get(sandbox, "77015550821", "status")[0] == "opted_out"
    assert "стоп" in page.lower()


def test_redraft_and_a_new_outcome_drop_the_old_text(client, sandbox: CrmConfig) -> None:  # noqa: F811
    contact = _draft(client, sandbox, "77015550831")
    client.post(f"/campaign/{contact}/message", data={"action": "redraft"})
    assert _get(sandbox, "77015550831", "status", "draft_text") == ("ready", None)
    assert "черновик пишется" in client.get("/campaign/?tab=drafts").get_data(as_text=True)

    _set(sandbox, "77015550831", status="draft", draft_text="Старый текст?")
    client.post(f"/campaign/{contact}/outcome", data={"outcome": "came"})
    assert _get(sandbox, "77015550831", "status", "segment", "draft_text") == ("ready", "after_trial", None)


def test_template_drafts_are_rewritten_on_request(client, sandbox: CrmConfig) -> None:  # noqa: F811
    _draft(client, sandbox, "77015550841")
    _draft(client, sandbox, "77015550842")
    _set(sandbox, "77015550841", draft_source="template")

    assert "Переписать моделью шаблонные (1)" in client.get("/campaign/?tab=drafts").get_data(as_text=True)
    client.post("/campaign/redraft-templates")

    assert _get(sandbox, "77015550841", "status")[0] == "ready"
    assert _get(sandbox, "77015550842", "status")[0] == "draft"


def test_an_empty_or_long_text_is_not_approved(client, sandbox: CrmConfig) -> None:  # noqa: F811
    contact = _draft(client, sandbox, "77015550851")
    assert "Пустое сообщение" in client.post(f"/campaign/{contact}/message", data={"action": "approve", "text": " "},
                                             follow_redirects=True).get_data(as_text=True)
    assert "Слишком длинно" in client.post(f"/campaign/{contact}/message",
                                           data={"action": "approve", "text": "а" * 1001},
                                           follow_redirects=True).get_data(as_text=True)
    assert _get(sandbox, "77015550851", "status")[0] == "draft"


def test_results_after_sending_and_the_funnel(client, sandbox: CrmConfig) -> None:  # noqa: F811
    contact = _draft(client, sandbox, "77015550861")
    now = datetime.now(tz=UTC).strftime("%Y-%m-%d %H:%M:%S.%f")
    _set(sandbox, "77015550861", status="replied", approved_text="Текст?", sent_at=now, replied_at=now)
    page = client.get("/campaign/?tab=sent").get_data(as_text=True)
    assert "Ответил" in page and "Текст?" in page

    client.post(f"/campaign/{contact}/result", data={"result": "bought", "note": "оплатили в зале"})

    assert _get(sandbox, "77015550861", "status", "note") == ("bought", "оплатили в зале")
    page = client.get("/campaign/").get_data(as_text=True)
    assert "из них купили абонемент: 1" in page and "Сегодня отправлено: 1" in page
    # После «стоп» итог не ставится.
    _set(sandbox, "77015550861", status="opted_out")
    page = client.post(f"/campaign/{contact}/result", data={"result": "booked"}, follow_redirects=True)
    assert "не после «стоп»" in page.get_data(as_text=True)


def test_the_switch_is_the_same_setting_as_in_settings(client, sandbox: CrmConfig) -> None:  # noqa: F811
    _upload(client, [_row("77015550871", "Интерес затих до записи", check="Нет")])
    page = client.get("/campaign/").get_data(as_text=True)
    assert "Включить рассылку" in page and "Бот выключен" in page

    page = client.post("/campaign/switch", data={"state": "on"}, follow_redirects=True).get_data(as_text=True)
    assert "бот выключен" in page
    store = AdminStore(sandbox.admin_db)
    try:
        assert store.get("campaign_enabled") == "on"
        store.set("bot_enabled", "on")
    finally:
        store.close()
    page = client.get("/campaign/").get_data(as_text=True)
    assert "Выключить рассылку" in page and "Бот выключен" not in page
    client.post("/campaign/switch", data={"state": "off"})
    store = AdminStore(sandbox.admin_db)
    try:
        assert store.get("campaign_enabled") == "off"
    finally:
        store.close()


def test_the_daily_limit_is_checked_in_settings(client, sandbox: CrmConfig) -> None:  # noqa: F811
    from tests.test_crm import _Form

    parser = _Form()
    parser.feed(client.get("/settings/").get_data(as_text=True))
    form = dict(parser.data)
    form["campaign_daily_limit"] = "500"
    page = client.post("/settings/", data=form, follow_redirects=True).get_data(as_text=True)
    assert "нужно число от 0 до 100" in page
    form["campaign_daily_limit"] = "20"
    client.post("/settings/", data=form)
    store = AdminStore(sandbox.admin_db)
    try:
        assert store.get("campaign_daily_limit") == "20"
    finally:
        store.close()


def test_the_old_tab_link_still_works(client, sandbox: CrmConfig) -> None:  # noqa: F811
    _draft(client, sandbox, "77015550881", "Старая закладка?")
    assert "Старая закладка?" in client.get("/campaign/?tab=ready").get_data(as_text=True)


def test_the_pult_warns_when_the_model_does_not_answer(client, sandbox: CrmConfig) -> None:  # noqa: F811
    """Кредиты Gemini кончились — включать рассылку рано: ответившие получат «сбой»."""
    _upload(client, [_row("77015550891", "Интерес затих до записи", check="Нет")])
    stamp = datetime.now(tz=UTC).strftime("%Y-%m-%d %H:%M:%S.%f")
    with closing(sqlite3.connect(sandbox.bot_db)) as conn, conn:
        columns = {row[1] for row in conn.execute("PRAGMA table_info(llm_call)")}
        for index in range(3):
            values = {"id": f"llm-{index}", "created_at": stamp, "error": "llm_quota: кредиты кончились"}
            for name, default in (("model", "m"), ("prompt_tokens", 0), ("candidates_tokens", 0), ("cost_usd", 0),
                                  ("latency_ms", 0), ("correlation_id", "x"), ("kb_hash", "h")):
                if name in columns:
                    values[name] = default
            names = ", ".join(values)
            conn.execute(f"INSERT INTO llm_call ({names}) VALUES ({', '.join('?' for _ in values)})", list(values.values()))

    page = client.get("/campaign/").get_data(as_text=True)

    assert "получат карточки вместо разговора" in page and "Billing" in page
