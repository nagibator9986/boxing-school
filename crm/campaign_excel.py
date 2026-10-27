"""Excel разбора старых переписок → строки для списка возврата.

Колонки ищутся по названию, а не по месту: файл собирает внешний разбор, и порядок
колонок в нём уже менялся. Берётся только то, что нужно для решения и первого
сообщения. Колонку «Фрагменты переписки» модуль не читает намеренно: цитаты из чужих
разговоров с именами детей хранить незачем, итога и препятствия достаточно.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from datetime import UTC, datetime, timedelta, timezone
from typing import IO, Any, Final

from app.campaign.contacts import CampaignSegment, segment_for

__all__ = ["IMPORTABLE_SHEETS", "ImportRow", "ParsedFile", "parse_workbook"]

#: Лист со списком к работе. Если его нет — первый лист с нужными колонками.
PREFERRED_SHEET: Final[str] = "Потенциальные"

#: Листы, которые можно загрузить, и что в них. «Проверить контекст» — переписки, которые
#: разбор не понял (голосовые, вложения, звонки): всё уходит на проверку человеку. Лист «Все
#: контакты» не загружается: там и действующие клиенты, и отказы, и сотрудники.
IMPORTABLE_SHEETS: Final[dict[str, str]] = {
    PREFERRED_SHEET: "потенциальные клиенты",
    "Проверить контекст": "разбор не понял переписку — всё на проверку",
}

#: Время в файле — по часам Костаная (колонки помечены «UTC+5»).
_FILE_ZONE: Final[timezone] = timezone(timedelta(hours=5))

#: Заголовок колонки → поле строки.
_COLUMNS: Final[dict[str, str]] = {
    "номер whatsapp": "phone",
    "имя из whatsapp": "wa_name",
    "имя родителя из текста": "parent_name",
    "имя ребенка / участника": "child_name",
    "этап и статус": "stage",
    "краткий итог переписки": "summary",
    "почему включен / исключен": "reason",
    "препятствие со слов клиента": "obstacle",
    "следующий шаг": "next_step",
    "нужна проверка": "check_required",
    "что проверить": "check_note",
    "последнее сообщение (utc+5)": "last_contact_at",
    "пробное / курс": "trial_info",
    "оплата по переписке": "payment_info",
    "связаться не раньше (utc+5)": "not_before",
    "статус первого анализа": "first_status",
    "ключ контакта": "source_key",
}

#: Без этих колонок файл не тот.
_REQUIRED: Final[frozenset[str]] = frozenset({"phone", "stage"})

#: Сколько строк сверху просматривается в поисках заголовка: над ним в файле сводка.
_HEADER_SEARCH_ROWS: Final[int] = 30

#: Длина текстовых полей в базе — с запасом для итога, коротко для остального.
_LIMITS: Final[dict[str, int]] = {"summary": 1500, "reason": 600, "obstacle": 600, "check_note": 600}
_DEFAULT_LIMIT: Final[int] = 300


@dataclass(frozen=True, slots=True)
class ImportRow:
    """Один контакт из файла."""

    chat_id: str
    segment: CampaignSegment
    stage: str | None
    check_required: bool
    wa_name: str | None = None
    parent_name: str | None = None
    child_name: str | None = None
    summary: str | None = None
    reason: str | None = None
    obstacle: str | None = None
    next_step: str | None = None
    check_note: str | None = None
    trial_info: str | None = None
    payment_info: str | None = None
    last_contact_at: datetime | None = None
    not_before: datetime | None = None
    source_key: str | None = None


@dataclass(slots=True)
class ParsedFile:
    """Итог чтения: строки, лист и что не разобралось — администратору это показывается."""

    sheet: str | None = None
    rows: list[ImportRow] = field(default_factory=list)
    errors: list[str] = field(default_factory=list)
    duplicates: int = 0


def parse_workbook(source: str | IO[bytes], *, sheet: str | None = None) -> ParsedFile:
    """Читает файл Excel. Ошибки формата — в ``errors``, а не исключением.

    ``sheet`` — какой лист читать; без него — «Потенциальные», а нет его — первый лист с
    нужными колонками.
    """
    from openpyxl import load_workbook

    result = ParsedFile()
    try:
        book = load_workbook(source, read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 - битый или не тот файл
        result.errors.append(f"Файл не читается как Excel: {type(exc).__name__}")
        return result
    try:
        if sheet is not None:
            sheets = [page for page in book.worksheets if page.title == sheet]
            if not sheets:
                result.errors.append(f"В файле нет листа «{sheet}».")
                return result
        else:
            sheets = sorted(book.worksheets, key=lambda page: page.title != PREFERRED_SHEET)
        for page in sheets:
            rows = page.iter_rows(values_only=True)
            header_at, columns = _find_header(rows)
            if columns is None:
                continue
            result.sheet = page.title
            seen: set[str] = set()
            for offset, values in enumerate(rows, start=header_at + 2):
                record = {name: values[index] for name, index in columns.items() if index < len(values)}
                if all(_clean(value) is None for value in record.values()):
                    continue
                row = _row(record, offset, result.errors)
                if row is None:
                    continue
                if row.chat_id in seen:
                    result.duplicates += 1
                    continue
                seen.add(row.chat_id)
                result.rows.append(row)
            return result
        result.errors.append(
            "В файле нет листа с колонками «Номер WhatsApp» и «Этап и статус»."
        )
        return result
    finally:
        book.close()


def _find_header(rows: Any) -> tuple[int, dict[str, int] | None]:
    for number, values in enumerate(rows):
        if number >= _HEADER_SEARCH_ROWS:
            break
        columns = {
            _COLUMNS[key]: index
            for index, value in enumerate(values)
            if (key := _header_key(value)) in _COLUMNS
        }
        if _REQUIRED <= columns.keys():
            return number, columns
    return 0, None


def _header_key(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "").lower().replace("ё", "е")).strip()


def _row(record: dict[str, Any], line: int, errors: list[str]) -> ImportRow | None:
    chat_id = _chat_id(record.get("phone"))
    if chat_id is None:
        errors.append(f"Строка {line}: номер не похож на казахстанский или российский — пропущена.")
        return None
    stage = _clean(record.get("stage"))
    return ImportRow(
        chat_id=chat_id,
        segment=segment_for(stage, _clean(record.get("first_status"))),
        stage=stage,
        check_required=_yes(record.get("check_required")),
        wa_name=_text(record, "wa_name"),
        parent_name=_text(record, "parent_name"),
        child_name=_text(record, "child_name"),
        summary=_text(record, "summary"),
        reason=_text(record, "reason"),
        obstacle=_text(record, "obstacle"),
        next_step=_text(record, "next_step"),
        check_note=_text(record, "check_note"),
        trial_info=_text(record, "trial_info"),
        payment_info=_text(record, "payment_info"),
        last_contact_at=_moment(record.get("last_contact_at")),
        not_before=_moment(record.get("not_before")),
        source_key=_text(record, "source_key"),
    )


def _chat_id(value: Any) -> str | None:
    """Номер как ``chat_id`` WhatsApp: 11 цифр с семёркой впереди."""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    digits = re.sub(r"\D", "", str(value or ""))
    if len(digits) == 11 and digits.startswith("8"):
        digits = "7" + digits[1:]
    elif len(digits) == 10:
        digits = "7" + digits
    return digits if len(digits) == 11 and digits.startswith("7") else None


def _clean(value: Any) -> str | None:
    text = re.sub(r"\s+", " ", str(value)).strip() if value is not None else ""
    return text or None


def _text(record: dict[str, Any], name: str) -> str | None:
    text = _clean(record.get(name))
    if text is None:
        return None
    limit = _LIMITS.get(name, _DEFAULT_LIMIT)
    return text if len(text) <= limit else text[: limit - 1].rstrip() + "…"


def _yes(value: Any) -> bool:
    """«Нужна проверка»: всё, кроме явного «Нет», — да. Сомнение решается в пользу проверки."""
    return (_clean(value) or "").lower() not in ("нет", "no", "false", "0")


def _moment(value: Any) -> datetime | None:
    """Время из файла (UTC+5) в UTC."""
    if isinstance(value, datetime):
        moment = value
    else:
        text = _clean(value)
        if text is None:
            return None
        moment = None
        for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%Y-%m-%d", "%d.%m.%Y %H:%M", "%d.%m.%Y"):
            try:
                moment = datetime.strptime(text, fmt)
                break
            except ValueError:
                continue
        if moment is None:
            return None
    if moment.tzinfo is None:
        moment = moment.replace(tzinfo=_FILE_ZONE)
    return moment.astimezone(UTC)
